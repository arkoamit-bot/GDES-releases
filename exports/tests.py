"""
Tests for the research-dataset export: structure, de-identification default,
0/1 boolean encoding, identified-export RBAC, and CSV/XLSX output.
"""
import datetime as dt

from django.contrib.auth.models import Group, User
from django.core.management import call_command
from rest_framework.authtoken.models import Token
from rest_framework.test import APITestCase

from analytics.services.outcomes import compute_patient_outcome
from labs.services.results import record_result
from patients.models import Patient
from treatments.models import DrugClass, DrugMaster, TreatmentExposure

from .services.dataset import build_dataset, columns
from .services.writers import to_csv, to_xlsx


def _patient():
    p = Patient.objects.create(
        patient_id="EXP-1", name="Secret Name", phone="0123", hospital_id="H9",
        sex="M", dob=dt.date(1970, 1, 1), enrollment_date=dt.date(2024, 1, 1),
        diabetes_status="t2", primary_diagnosis="IgA nephropathy")
    record_result(p, "creatinine", result_date=dt.date(2024, 1, 1), value_numeric=1.2)
    dapa = DrugMaster.objects.create(generic_name="Dapagliflozin", drug_class=DrugClass.SGLT2I)
    TreatmentExposure.objects.create(patient=p, drug=dapa, drug_name="Dapagliflozin",
                                     start_date=dt.date(2024, 1, 1), ongoing=True)
    compute_patient_outcome(p)
    return p


class DatasetStructureTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_labs", verbosity=0)

    def test_deidentified_by_default_excludes_pii(self):
        _patient()
        cols, rows = build_dataset(Patient.objects.all())
        self.assertNotIn("name", cols)
        self.assertNotIn("phone", cols)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["patient_id"], "EXP-1")
        self.assertTrue(rows[0]["ever_sglt2i"])      # exposure flag populated

    def test_identified_includes_pii(self):
        _patient()
        cols, rows = build_dataset(Patient.objects.all(), identified=True)
        self.assertIn("name", cols)
        self.assertEqual(rows[0]["name"], "Secret Name")

    def test_csv_encodes_booleans_as_0_1(self):
        _patient()
        cols, rows = build_dataset(Patient.objects.all())
        csv_text = to_csv(cols, rows)
        header = csv_text.splitlines()[0]
        self.assertIn("ever_sglt2i", header)
        # ever_sglt2i True -> "1" somewhere in the data row.
        self.assertIn("1", csv_text.splitlines()[1])

    def test_xlsx_bytes_produced(self):
        _patient()
        cols, rows = build_dataset(Patient.objects.all())
        data = to_xlsx(cols, rows)
        self.assertTrue(data[:2] == b"PK")           # xlsx is a zip

    def test_data_dictionary_covers_every_column(self):
        from .services.dictionary import data_dictionary
        cols = columns()
        entries = data_dictionary()
        self.assertEqual([e["column"] for e in entries], cols)   # same order, full coverage
        # Every column has a description and type.
        for e in entries:
            self.assertTrue(e["description"], f"missing description for {e['column']}")
            self.assertTrue(e["type"], f"missing type for {e['column']}")

    def test_xlsx_includes_data_dictionary_sheet(self):
        import io
        from openpyxl import load_workbook
        from .services.dictionary import DICTIONARY_COLUMNS, data_dictionary
        _patient()
        cols, rows = build_dataset(Patient.objects.all())
        data = to_xlsx(cols, rows, dictionary=data_dictionary(),
                       dictionary_columns=DICTIONARY_COLUMNS)
        wb = load_workbook(io.BytesIO(data))
        self.assertIn("data_dictionary", wb.sheetnames)
        self.assertIn("research_dataset", wb.sheetnames)


class ExportRBACTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_labs", verbosity=0)
        call_command("seed_roles")

    def _auth(self, username, role):
        u = User.objects.create_user(username, password="x")
        u.groups.add(Group.objects.get(name=role))
        token, _ = Token.objects.get_or_create(user=u)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")

    def test_unauthenticated_blocked(self):
        self.assertEqual(self.client.get("/exports/research-dataset/").status_code, 401)

    def test_readonly_can_export_deidentified(self):
        _patient()
        self._auth("ro", "readonly")
        resp = self.client.get("/exports/research-dataset/?fmt=csv")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("text/csv", resp["Content-Type"])
        self.assertNotIn(b"Secret Name", resp.content)

    def test_readonly_cannot_export_identified(self):
        _patient()
        self._auth("ro2", "readonly")
        resp = self.client.get("/exports/research-dataset/?identified=1")
        self.assertEqual(resp.status_code, 403)

    def test_data_manager_can_export_identified(self):
        _patient()
        self._auth("dm", "data_manager")
        resp = self.client.get("/exports/research-dataset/?identified=1")
        self.assertEqual(resp.status_code, 200)
        self.assertIn(b"Secret Name", resp.content)
